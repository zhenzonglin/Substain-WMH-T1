"""Prepare the healthy LS cohort without modifying GB or source data.

``prepare_inputs(config)`` writes only in ``ls.preparation_dir`` and the
explicit ``participants_LS.tsv`` destination. Expected input errors are
collected before the reports are written; subjects missing either image are
recorded and excluded. Other input errors prevent participant publication.
Normative ineligibility is reported, not an input exclusion.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

import nibabel as nib
import numpy as np


REQUIRED_HEADERS = ("BIGDATA_CODE", "AGE", "SEX")
PARTICIPANT_COLUMNS = ("participant_id", "age", "sex", "site_id", "t1w", "flair", "lesion_mask")
METADATA_COLUMNS = PARTICIPANT_COLUMNS[:4]
ID_PATTERN = re.compile(r"P\d+", re.ASCII)
MASK_NAME = "healthy_zero_lesion_MNI152_1mm_LS.nii.gz"


class InputPreparationError(ValueError):
    """Input preparation failed; the report contains every detected issue."""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _absolute(root: Path, value: object) -> Path:
    path = Path(str(value)).expanduser()
    return (path if path.is_absolute() else root / path).resolve()


def _inside(path: Path, directory: Path) -> bool:
    try:
        path.relative_to(directory)
        return True
    except ValueError:
        return False


def _destination(path: Path, allowed: Sequence[Path]) -> Path:
    if path.is_symlink() or not any(_inside(path.resolve(), base) for base in allowed):
        raise InputPreparationError("Unsafe LS output destination: {}".format(path))
    return path


def _atomic_text(path: Path, text: str, allowed: Sequence[Path]) -> bool:
    _destination(path, allowed)
    data = text.encode("utf-8")
    if path.is_file() and path.read_bytes() == data:
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=str(path.parent), prefix=".LS-", delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(data)
        os.replace(str(temporary), str(path))
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return True


def _tsv(columns: Sequence[str], rows: Sequence[Mapping[str, object]]) -> str:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=list(columns), delimiter="\t", lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue()


def _json(payload: object) -> str:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n"


def _issue(issues: List[dict], code: str, message: str, participant_id: str = "", row: object = "", severity: str = "error") -> None:
    issues.append({"severity": severity, "code": code, "participant_id": participant_id, "source_row": row, "message": message})


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _cell_text(value: object) -> str:
    return "" if value is None else str(value)


def _number(value: object, label: str) -> float:
    if isinstance(value, bool) or _blank(value):
        raise ValueError("{} must be a finite number".format(label))
    number = float(value)
    if not math.isfinite(number):
        raise ValueError("{} must be a finite number".format(label))
    return number


def _read_clinical(path: Path, issues: List[dict]) -> Tuple[List[str], List[dict], Dict[str, dict], dict, List[dict]]:
    try:
        import openpyxl
    except ImportError as exc:
        _issue(issues, "xlsx_dependency_missing", "openpyxl is required in the isolated LS dependency path: {}".format(exc))
        return [], [], {}, {}, []
    workbook = cached = None
    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
        cached = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        candidates = []
        for sheet in workbook.worksheets:
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = [str(value).lstrip("\ufeff").strip() if value is not None else "" for value in first]
            if all(name in headers for name in REQUIRED_HEADERS):
                candidates.append((sheet.title, first, headers))
        if len(candidates) != 1:
            _issue(issues, "clinical_sheet_count", "Expected one sheet with first-row BIGDATA_CODE/AGE/SEX; found {}: {}".format(len(candidates), [item[0] for item in candidates]))
            return [], [], {}, {"candidate_sheets": [item[0] for item in candidates]}, []
        title, raw_first, headers = candidates[0]
        duplicates = sorted(name for name, count in Counter(headers).items() if name and count > 1)
        if duplicates:
            _issue(issues, "duplicate_clinical_headers", "Duplicate normalized headers in {}: {}".format(title, duplicates))
            return [], [], {}, {"sheet": title, "header_row": 1}, []
        # Blank columns without data are harmless; columns with data need a name.
        rows = list(workbook[title].iter_rows(min_row=2))
        cached_rows = list(cached[title].iter_rows(min_row=2, values_only=True))
        used = [index for index, name in enumerate(headers) if name or any(index < len(row) and not _blank(row[index].value) for row in rows)]
        if any(not headers[index] for index in used):
            _issue(issues, "unnamed_clinical_column", "A populated column has no first-row header in {}".format(title))
            return [], [], {}, {"sheet": title, "header_row": 1}, []
        original_headers = [str(raw_first[index]) for index in used]
        raw_records: List[dict] = []
        records: Dict[str, dict] = {}
        formulas: List[dict] = []
        id_rows: Dict[str, List[int]] = {}
        for row_number, cells in enumerate(rows, start=2):
            if all(_blank(cell.value) for cell in cells):
                continue
            values = []
            cache = cached_rows[row_number - 2] if row_number - 2 < len(cached_rows) else ()
            for index in range(len(headers)):
                cell = cells[index] if index < len(cells) else None
                value = cell.value if cell is not None else None
                if cell is not None and cell.data_type == "f":
                    value = cache[index] if index < len(cache) else None
                    formulas.append({"sheet": title, "cell": cell.coordinate, "formula": cell.value, "cached_value": _cell_text(value)})
                    if value is None:
                        _issue(issues, "formula_cache_missing", "{}!{} has no cached formula value".format(title, cell.coordinate), row=row_number)
                if cell is not None and cell.data_type == "e":
                    _issue(issues, "excel_cell_error", "{}!{} contains Excel error {}".format(title, cell.coordinate, cell.value), row=row_number)
                values.append(value)
            raw_records.append({str(raw_first[index]): _cell_text(values[index]) for index in used})
            normalized = dict(zip(headers, values))
            raw_id = normalized.get("BIGDATA_CODE")
            participant = raw_id.strip() if isinstance(raw_id, str) else ""
            if not ID_PATTERN.fullmatch(participant):
                _issue(issues, "invalid_clinical_id", "BIGDATA_CODE must be a P-prefixed string preserving digits: {!r}".format(raw_id), row=row_number)
                continue
            id_rows.setdefault(participant, []).append(row_number)
            record: Dict[str, Any] = {"participant_id": participant, "source_row": row_number, "clinical_valid": True}
            try:
                age = _number(normalized.get("AGE"), "AGE")
                if not 0 < age <= 120:
                    raise ValueError("AGE must be in (0,120]")
                record["age"] = age
            except (TypeError, ValueError) as exc:
                record["clinical_valid"] = False
                _issue(issues, "invalid_age", str(exc), participant, row_number)
            try:
                sex_code = _number(normalized.get("SEX"), "SEX")
                if sex_code not in (1.0, 2.0):
                    raise ValueError("SEX must be exactly 1 (male) or 2 (female)")
                record["sex"] = {1.0: "male", 2.0: "female"}[sex_code]
            except (TypeError, ValueError) as exc:
                record["clinical_valid"] = False
                _issue(issues, "invalid_sex", str(exc), participant, row_number)
            records.setdefault(participant, record)
        for participant, locations in sorted(id_rows.items()):
            if len(locations) > 1:
                _issue(issues, "duplicate_clinical_id", "Duplicate BIGDATA_CODE at rows {}".format(locations), participant)
                records[participant]["clinical_valid"] = False
        return original_headers, raw_records, records, {"sheet": title, "header_row": 1, "source_row_count": len(raw_records)}, formulas
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        _issue(issues, "clinical_read_error", "{}: {}".format(type(exc).__name__, exc))
        return [], [], {}, {}, []
    finally:
        if workbook is not None:
            workbook.close()
        if cached is not None:
            cached.close()


def _scan_images(root: Path, issues: List[dict]) -> Dict[str, dict]:
    result: Dict[str, dict] = {}
    if not root.is_dir():
        _issue(issues, "input_root_missing", "Input directory does not exist: {}".format(root))
        return result
    for directory in sorted(root.iterdir()):
        if not directory.is_dir() or not directory.name.startswith("P"):
            continue
        participant = directory.name
        if not ID_PATTERN.fullmatch(participant):
            _issue(issues, "invalid_image_id", "Invalid P-prefixed subject directory: {}".format(directory), participant)
            continue
        record: Dict[str, Any] = {"participant_id": participant, "image_directory": str(directory.resolve())}
        for modality, pattern in (("t1w", "*T1w*.nii.gz"), ("flair", "*T2wFLAIR*.nii.gz")):
            paths = sorted(path.resolve() for path in directory.glob(pattern) if path.is_file())
            record[modality + "_candidates"] = [str(path) for path in paths]
            record[modality] = str(paths[0]) if len(paths) == 1 else ""
            if len(paths) != 1:
                _issue(issues, "{}_count".format(modality), "Expected one {} file, found {}: {}".format(modality, len(paths), paths), participant, severity="warning" if not paths else "error")
            else:
                try:
                    nib.load(str(paths[0]))
                except (OSError, ValueError, nib.filebasedimages.ImageFileError) as exc:
                    _issue(issues, "invalid_nifti", "{}: {}".format(paths[0], exc), participant)
        result[participant] = record
    if not result:
        _issue(issues, "no_image_subjects", "No direct P<digits> subject directories found")
    return result


def _reference(config: Mapping[str, object], root: Path, issues: List[dict]) -> Tuple[List[dict], dict]:
    try:
        t1 = config["t1"]
        path = _absolute(root, t1["genmind_csv"])
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not {"Age", "Sex"}.issubset(reader.fieldnames or []):
                raise ValueError("GenMIND reference must contain Age and Sex")
            rows = []
            for number, row in enumerate(reader, start=2):
                age = _number(row["Age"], "GenMIND Age row {}".format(number))
                sex = {"F": "female", "M": "male", "Female": "female", "Male": "male"}.get(row["Sex"])
                if sex is None:
                    raise ValueError("Unknown GenMIND Sex at row {}: {!r}".format(number, row["Sex"]))
                rows.append({"age": age, "sex": sex})
        minimum = int(t1.get("min_reference_n", 300))
        narrow = _number(t1.get("narrow_age_half_window", 2.5), "narrow age window")
        expanded = _number(t1.get("expanded_age_half_window", 5.0), "expanded age window")
        if minimum < 1 or narrow <= 0 or expanded < narrow:
            raise ValueError("Invalid inherited GenMIND count or age window")
        return rows, {"path": str(path), "sha256": sha256(path), "row_count": len(rows), "min_reference_n": minimum, "narrow_age_half_window": narrow, "expanded_age_half_window": expanded, "scope": "age_sex_reference_count_only"}
    except (KeyError, TypeError, ValueError, OSError) as exc:
        _issue(issues, "normative_reference_error", "{}: {}".format(type(exc).__name__, exc))
        return [], {}


def _eligibility(clinical: Mapping[str, dict], references: Sequence[dict], settings: Mapping[str, object], image_ids: Sequence[str]) -> List[dict]:
    result = []
    for participant in sorted(image_ids):
        source = clinical.get(participant, {})
        row: Dict[str, Any] = {"participant_id": participant, "age": source.get("age", ""), "sex": source.get("sex", ""), "eligible": False, "reference_n": 0, "age_half_window": "", "failure_reason": "clinical_missing_or_invalid", "scope": "age_sex_reference_count_only"}
        if source.get("clinical_valid"):
            age, sex = source["age"], source["sex"]
            if not 22 <= age <= 90:
                row["failure_reason"] = "age_outside_22_90"
            elif not settings:
                row["failure_reason"] = "normative_reference_unavailable"
            else:
                minimum = int(settings["min_reference_n"])
                for window in (float(settings["narrow_age_half_window"]), float(settings["expanded_age_half_window"])):
                    count = sum(item["sex"] == sex and age - window <= item["age"] <= age + window for item in references)
                    row.update(reference_n=count, age_half_window=window)
                    if count >= minimum:
                        break
                row.update(eligible=count >= minimum, failure_reason="" if count >= minimum else "reference_n_below_{}".format(minimum))
        result.append(row)
    return result


def _template(path: Path, issues: List[dict]) -> Tuple[Optional[Any], dict]:
    try:
        image = nib.load(str(path))
        if len(image.shape) != 3 or not np.allclose(image.header.get_zooms()[:3], (1.0, 1.0, 1.0), atol=1e-5):
            raise ValueError("LS zero-mask template must be a 3D 1 mm FSL reference")
        if not np.isfinite(image.affine).all() or abs(np.linalg.det(image.affine[:3, :3])) < 1e-8:
            raise ValueError("Template has an invalid affine")
        return image, {"path": str(path), "sha256": sha256(path), "shape": list(image.shape), "affine": image.affine.tolist()}
    except (OSError, ValueError, nib.filebasedimages.ImageFileError) as exc:
        _issue(issues, "template_error", "{}: {}".format(type(exc).__name__, exc))
        return None, {}


def _zero_mask(path: Path, template: Any, allowed: Sequence[Path]) -> None:
    _destination(path, allowed)
    if path.exists():
        image = nib.load(str(path))
        qform, qcode = image.get_qform(coded=True)
        sform, scode = image.get_sform(coded=True)
        if (image.shape != template.shape or not np.allclose(image.affine, template.affine, atol=1e-5)
                or np.count_nonzero(np.asanyarray(image.dataobj)) != 0
                or image.get_data_dtype() != np.dtype(np.uint8)
                or not int(qcode) or not int(scode) or not np.allclose(qform, sform, atol=1e-4)):
            raise InputPreparationError("Existing synthetic zero mask differs from its contract: {}".format(path))
        return
    header = template.header.copy()
    header.set_data_dtype(np.uint8)
    header.set_slope_inter(1.0, 0.0)
    image = nib.Nifti1Image(np.zeros(template.shape, dtype=np.uint8), template.affine, header)
    image.set_qform(template.affine, code=1)
    image.set_sform(template.affine, code=1)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=str(path.parent), prefix=".zero_LS-", suffix=".nii.gz", delete=False) as handle:
            temporary = Path(handle.name)
        nib.save(image, str(temporary))
        os.replace(str(temporary), str(path))
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def prepare_inputs(config: dict) -> dict:
    """Prepare complete image pairs, preserving normative-ineligible records.

    ``config['ls']`` requires ``input_root``, ``clinical_file``,
    ``preparation_dir`` and ``template``; ``site_id`` defaults to ``LS``.
    The inherited ``t1`` GenMIND settings are required. Return ``status`` and
    paths/counts. The caller must refuse a run when ``normative_eligible_count``
    is zero and must retain the input report for the run it launches.
    """
    root = Path(str(config["project_root"])).resolve()
    settings = config["ls"]
    source_root = _absolute(root, settings["input_root"])
    clinical_path = _absolute(root, settings["clinical_file"])
    raw_prep = Path(str(settings["preparation_dir"])).expanduser()
    raw_prep = raw_prep if raw_prep.is_absolute() else root / raw_prep
    if raw_prep.is_symlink():
        raise InputPreparationError("Preparation directory must not be a symlink")
    prep = raw_prep.resolve()
    participants = Path(str(config["participants"])).expanduser()
    if not participants.is_absolute() or participants.name != "participants_LS.tsv":
        raise InputPreparationError("participants must be an absolute participants_LS.tsv path")
    if participants.parent.resolve() != (root / "config").resolve():
        raise InputPreparationError("participants_LS.tsv must be in the existing project's config directory")
    if _inside(prep, source_root) or _inside(source_root, prep) or _inside(clinical_path, prep):
        raise InputPreparationError("LS preparation and source paths must not overlap")
    allowed = (prep, participants.parent.resolve())
    _destination(participants, allowed)
    issues: List[dict] = []
    source_fingerprint: Dict[str, object] = {"path": str(clinical_path)}
    try:
        source_fingerprint.update(sha256=sha256(clinical_path), size_bytes=clinical_path.stat().st_size)
    except OSError as exc:
        _issue(issues, "clinical_file_missing", str(exc))
    headers, raw_clinical, clinical, sheet, formulas = _read_clinical(clinical_path, issues)
    if "sha256" in source_fingerprint:
        try:
            if sha256(clinical_path) != source_fingerprint["sha256"]:
                _issue(issues, "clinical_changed_during_read", "LS_INFO.xlsx changed while it was being read; retry with a stable source file")
        except OSError as exc:
            _issue(issues, "clinical_changed_during_read", str(exc))
    images = _scan_images(source_root, issues)
    references, reference_info = _reference(config, root, issues)
    template_path = _absolute(root, settings["template"])
    template, template_info = _template(template_path, issues)
    site = str(settings.get("site_id", "LS")).strip()
    if not site:
        _issue(issues, "invalid_site_id", "site_id must not be blank")
    complete_image_ids = sorted(participant for participant, image in images.items()
                                if all(len(image[modality + "_candidates"]) == 1 for modality in ("t1w", "flair")))
    if not complete_image_ids:
        _issue(issues, "no_complete_image_pairs", "No subjects have exactly one matched T1w and one matched T2wFLAIR image; participants were not published")
    matching = []
    excluded = []
    manifest = []
    for participant in sorted(set(images) | set(clinical)):
        image = images.get(participant, {})
        row = clinical.get(participant, {})
        if participant in images and participant not in clinical:
            _issue(issues, "clinical_match_missing", "No matching BIGDATA_CODE record", participant)
        elif participant in clinical and participant not in images:
            _issue(issues, "clinical_only", "Clinical row retained; no image subject directory", participant, row.get("source_row", ""), "warning")
        missing = [modality for modality in ("t1w", "flair") if participant in images and not image[modality + "_candidates"]]
        reasons = ["missing_{}".format(modality) for modality in missing]
        if participant not in images:
            reasons.append("no_image_directory")
        else:
            reasons.extend("ambiguous_{}".format(modality) for modality in ("t1w", "flair") if len(image[modality + "_candidates"]) > 1)
        if not row.get("clinical_valid"):
            reasons.append("clinical_missing_or_invalid")
        if any(issue["code"] == "invalid_nifti" and issue["participant_id"] == participant for issue in issues):
            reasons.append("invalid_nifti")
        match = {"participant_id": participant, "image_present": participant in images, "clinical_present": participant in clinical, "clinical_valid": row.get("clinical_valid", False), "source_row": row.get("source_row", ""), "age": row.get("age", ""), "sex": row.get("sex", ""), "site_id": site, "image_directory": image.get("image_directory", ""), "t1w": image.get("t1w", ""), "flair": image.get("flair", ""), "t1w_candidates": json.dumps(image.get("t1w_candidates", []), ensure_ascii=False), "flair_candidates": json.dumps(image.get("flair_candidates", []), ensure_ascii=False), "included_for_analysis": not reasons, "exclusion_reason": ";".join(reasons)}
        matching.append(match)
        if missing:
            excluded.append(dict(match, missing_modalities=";".join(missing)))
        for modality in ("t1w", "flair"):
            for path_text in image.get(modality + "_candidates", []):
                path = Path(path_text)
                try:
                    manifest.append({"participant_id": participant, "modality": modality, "source": str(path), "source_sha256": sha256(path), "source_kind": "original_image"})
                except OSError as exc:
                    _issue(issues, "image_hash_error", "{}: {}".format(path, exc), participant)
    included_ids = [row["participant_id"] for row in matching if row["included_for_analysis"]]
    eligibility = _eligibility(clinical, references, reference_info, included_ids)
    for modality, fingerprint in (("clinical_xlsx", source_fingerprint), ("FSL_MNI152_template", template_info), ("GenMIND_reference", reference_info)):
        if "sha256" in fingerprint:
            manifest.append({"participant_id": "", "modality": modality, "source": fingerprint["path"], "source_sha256": fingerprint["sha256"], "source_kind": "original_source"})
    mask_path = prep / MASK_NAME
    if template is not None and mask_path.exists():
        try:
            _zero_mask(mask_path, template, allowed)
        except (OSError, ValueError, nib.filebasedimages.ImageFileError) as exc:
            _issue(issues, "existing_zero_mask_error", str(exc))
    report_path = prep / "input_report_LS.json"
    report: Dict[str, Any] = {"status": "fail" if any(item["severity"] == "error" for item in issues) else "pass", "cohort": "LS", "input_root": str(source_root), "clinical_source": dict(source_fingerprint, **sheet), "image_subject_count": len(images), "clinical_row_count": len(raw_clinical), "clinical_unique_id_count": len(clinical), "clinical_only_count": len(set(clinical) - set(images)), "participants_count": len(included_ids), "excluded_subject_count": len(excluded), "missing_t1w_count": sum(not image["t1w_candidates"] for image in images.values()), "missing_flair_count": sum(not image["flair_candidates"] for image in images.values()), "missing_modality_count_scope": "image_subject_directories; per-subject modality counts may overlap", "excluded_subjects_tsv": str(prep / "excluded_subjects_LS.tsv"), "normative_eligible_count": sum(bool(row["eligible"]) for row in eligibility), "normative_ineligible_count": sum(not bool(row["eligible"]) for row in eligibility), "normative_scope": "age_sex_reference_count_only; image-dependent normative checks remain in the original pipeline", "normative_reference": reference_info, "template": template_info, "healthy_no_acute_lesion": True, "lesion_source_kind": "synthetic_zero_mask", "participants_tsv": str(participants), "metadata_tsv": str(prep / "metadata_LS.tsv"), "clinical_tsv": str(prep / "clinical_LS.tsv"), "normative_eligibility_tsv": str(prep / "normative_eligibility_LS.tsv"), "input_manifest": str(prep / "input_manifest_LS.tsv"), "report_path": str(report_path), "raw_inputs_modified": False, "issues": issues}
    def write_reports() -> None:
        _atomic_text(prep / "clinical_LS.tsv", _tsv(headers, raw_clinical) if headers else "", allowed)
        _atomic_text(prep / "clinical_formulas_LS.tsv", _tsv(("sheet", "cell", "formula", "cached_value"), formulas), allowed)
        _atomic_text(prep / "input_matching_LS.tsv", _tsv(("participant_id", "image_present", "clinical_present", "clinical_valid", "source_row", "age", "sex", "site_id", "image_directory", "t1w", "flair", "t1w_candidates", "flair_candidates", "included_for_analysis", "exclusion_reason"), matching), allowed)
        _atomic_text(prep / "excluded_subjects_LS.tsv", _tsv(("participant_id", "missing_modalities", "exclusion_reason", "image_directory", "t1w", "flair", "t1w_candidates", "flair_candidates", "clinical_present", "clinical_valid", "source_row", "age", "sex", "site_id"), excluded), allowed)
        _atomic_text(prep / "input_issues_LS.tsv", _tsv(("severity", "code", "participant_id", "source_row", "message"), issues), allowed)
        _atomic_text(prep / "normative_eligibility_LS.tsv", _tsv(("participant_id", "age", "sex", "eligible", "reference_n", "age_half_window", "failure_reason", "scope"), eligibility), allowed)
        _atomic_text(prep / "input_manifest_LS.tsv", _tsv(("participant_id", "modality", "source", "source_sha256", "source_kind"), manifest), allowed)
        _atomic_text(report_path, _json(report), allowed)
    if report["status"] != "pass":
        write_reports()
        raise InputPreparationError("LS input preparation failed with {} errors; see {}".format(sum(item["severity"] == "error" for item in issues), report_path))
    _zero_mask(mask_path, template, allowed)
    mask_provenance = {"source_kind": "synthetic_zero_mask", "reason": "user_confirmed_healthy_no_acute_lesion", "observed_lesion_segmentation": False, "mask_path": str(mask_path), "mask_sha256": sha256(mask_path), "nonzero_voxel_count": 0, "template": template_info}
    report["synthetic_lesion"] = mask_provenance
    manifest.append({"participant_id": "", "modality": "healthy_zero_lesion", "source": str(mask_path), "source_sha256": mask_provenance["mask_sha256"], "source_kind": "synthetic_zero_mask"})
    rows = [{"participant_id": participant, "age": clinical[participant]["age"], "sex": clinical[participant]["sex"], "site_id": site, "t1w": images[participant]["t1w"], "flair": images[participant]["flair"], "lesion_mask": str(mask_path)} for participant in included_ids]
    _atomic_text(prep / "zero_mask_provenance_LS.json", _json(mask_provenance), allowed)
    _atomic_text(prep / "metadata_LS.tsv", _tsv(METADATA_COLUMNS, rows), allowed)
    write_reports()
    # Publish participants last, after the explicit exclusions and validation.
    _atomic_text(participants, _tsv(PARTICIPANT_COLUMNS, rows), allowed)
    print("LS prepare complete: image_subject_count={} participants_count={} excluded_subject_count={} excluded_subjects_tsv={}".format(
        report["image_subject_count"], report["participants_count"], report["excluded_subject_count"], report["excluded_subjects_tsv"]), flush=True)
    return report
